from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("Техническое_задание_портал_КНБ_2026-07-13_полная_версия.docx")

BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(0, 0, 0)
MUTED = RGBColor(89, 89, 89)
FILL = "EEF3F8"


def set_font(run, size=11, bold=False, italic=False, color=DARK):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def shade(cell, fill=FILL):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    if shd.getparent() is None:
        tc_pr.append(shd)


def cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    mar = tbl_pr.find(qn("w:tblCellMar")) or OxmlElement("w:tblCellMar")
    if mar.getparent() is None:
        tbl_pr.append(mar)
    for key, val in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = mar.find(qn(f"w:{key}")) or OxmlElement(f"w:{key}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        if node.getparent() is None:
            mar.append(node)


def set_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW")) or OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    if tbl_w.getparent() is None:
        tbl_pr.append(tbl_w)
    tbl_ind = tbl_pr.find(qn("w:tblInd")) or OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    if tbl_ind.getparent() is None:
        tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW")) or OxmlElement("w:tcW")
            tc_w.set(qn("w:w"), str(widths[i]))
            tc_w.set(qn("w:type"), "dxa")
            if tc_w.getparent() is None:
                tc_pr.append(tc_w)


def p(doc, text="", *, size=11, bold=False, italic=False, color=DARK, align=None, before=0, after=6, style=None):
    para = doc.add_paragraph(style=style)
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.10
    if align is not None:
        para.alignment = align
    if text:
        run = para.add_run(text)
        set_font(run, size=size, bold=bold, italic=italic, color=color)
    return para


def h(doc, text, level=1):
    para = doc.add_paragraph(style=f"Heading {level}")
    para.paragraph_format.keep_with_next = True
    para.paragraph_format.space_before = Pt({1: 16, 2: 11, 3: 7}[level])
    para.paragraph_format.space_after = Pt({1: 8, 2: 5, 3: 3}[level])
    run = para.add_run(text)
    set_font(run, size={1: 16, 2: 13, 3: 12}[level], bold=True, color=BLUE)
    return para


def bullets(doc, items):
    for item in items:
        para = p(doc, style="List Bullet", after=3)
        run = para.add_run(item)
        set_font(run)


def numbers(doc, items):
    for item in items:
        para = p(doc, style="List Number", after=3)
        run = para.add_run(item)
        set_font(run)


def kv(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    cell_margins(table)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        shade(cells[0])
        for idx, cell in enumerate(cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                for run in para.runs:
                    set_font(run, size=10, bold=(idx == 0))
    set_widths(table, [2300, 7060])
    p(doc, "", after=2)


def matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    cell_margins(table)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        shade(cell)
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                set_font(run, size=9.5, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for para in cells[i].paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    set_font(run, size=9.2)
    set_widths(table, widths)
    p(doc, "", after=2)


def configure(doc):
    s = doc.sections[0]
    s.page_width = Inches(8.5)
    s.page_height = Inches(11)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
    s.header_distance = s.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[name].font.name = "Calibri"
        styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        styles[name].font.bold = True
        styles[name].font.color.rgb = BLUE
    header = s.header.paragraphs[0]
    header.text = ""
    run = header.add_run("Техническое задание | Официальный интернет-портал КНБ РК")
    set_font(run, size=9, color=MUTED)
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Версия 2.0 | 13.07.2026")
    set_font(run, size=9, color=MUTED)


def section(doc, number, title):
    p(doc, f"Раздел {number} из 12", size=10, color=MUTED, after=1)
    h(doc, title, 1)


def title_page(doc):
    p(doc, "Техническое задание", size=24, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    p(
        doc,
        "на создание и развитие автоматизированной информационной системы официального интернет-портала Комитета национальной безопасности Республики Казахстан",
        size=15,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=16,
    )
    p(doc, "Раздел 1 из 12", color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    kv(
        doc,
        [
            ("Полное наименование", "Автоматизированная информационная система «Официальный интернет-портал Комитета национальной безопасности Республики Казахстан»."),
            ("Краткое наименование", "Портал КНБ РК."),
            ("Заказчик", "Комитет национальной безопасности Республики Казахстан."),
            ("Разработчик", "Проектная команда разработки портала КНБ РК, выполняющая frontend, backend, database, DevOps, testing и document delivery."),
            ("Основание версии", "Полная редакция ТЗ от 13.07.2026, подготовленная для закрытия функциональных, технических, организационных и приемочных вопросов без незаполненных разделов."),
            ("Статус", "Окончательная проектная редакция для разработки, внедрения, проверки и приемки системы."),
        ],
    )
    h(doc, "Аннотация", 1)
    p(doc, "Настоящее техническое задание фиксирует полный состав требований к порталу КНБ РК как к продуктивному государственному интернет-ресурсу. Документ объединяет позицию заказчика, автора ТЗ и разработчика: решения, сроки, состав функций, роли, данные, безопасность, эксплуатация и приемка сформулированы как конкретные обязательства проекта.")
    p(doc, "Основная общественная функция портала - официальное доведение до населения сведений о деятельности Комитета, направлениях его работы, пресс-материалах, контактной информации, порядке обращения, правилах поступления на службу и обучение, а также предоставление гражданам понятных цифровых каналов первичного взаимодействия с Комитетом.")
    p(doc, "ТЗ построено на текущем проектном репозитории, открытом официальном ресурсе КНБ РК на GOV.KZ и типовой структуре государственных интернет-ресурсов Республики Казахстан. Внешние сервисы eGov, e-Otinish и eLicense используются только как справочные внешние ссылки в публичном контенте; техническая интеграция с ними в данной версии не выполняется.", bold=True)
    doc.add_page_break()


def build():
    doc = Document()
    configure(doc)
    title_page(doc)

    section(doc, 2, "Назначение и цели создания системы")
    h(doc, "2.1. Шифр темы или номер договора", 2)
    p(doc, "Шифр проекта: KNB-PORTAL-2026. Номер внутреннего проектного дела: KNB-PORTAL-2026-TZ-01. Дата начала проектной реализации: 15.07.2026. Дата готовности первой продуктивной версии: 30.09.2026. Дата завершения гарантийного периода первой версии: 31.12.2026.")
    h(doc, "2.2. Заказчик, разработчик и пользователи", 2)
    kv(doc, [
        ("Заказчик", "Комитет национальной безопасности Республики Казахстан, 010000, Республика Казахстан, г. Астана, ул. Сыганак, 66."),
        ("Разработчик", "Проектная команда портала КНБ РК: frontend-разработчик, backend-разработчик, инженер БД, инженер безопасности, DevOps-инженер, тестировщик и технический писатель."),
        ("Публичные пользователи", "Граждане, кандидаты на службу или обучение, представители СМИ, посетители, получающие сведения о деятельности Комитета."),
        ("Служебные пользователи", "Администратор, модератор, ответственный сотрудник кадрового направления и технический оператор."),
    ])
    h(doc, "2.3. Документы и источники основания", 2)
    bullets(doc, [
        "Закон Республики Казахстан «Об органах национальной безопасности Республики Казахстан».",
        "Законодательство Республики Казахстан об информатизации, персональных данных, государственных секретах, обращениях физических и юридических лиц и официальных интернет-ресурсах.",
        "Официальный интернет-ресурс КНБ РК на GOV.KZ, включая разделы о Комитете, пресс-центр, направления деятельности, контакты, документы, прием на службу и работу с населением.",
        "ГОСТ 34.601 и структура исходного шаблона технического задания.",
        "Репозиторий проекта: apps/web, apps/api, docs/ARCHITECTURE.md, SECURITY_ROADMAP.md, docker-compose.yml и миграции Alembic.",
    ])
    h(doc, "2.4. Плановые сроки", 2)
    matrix(doc, ["Этап", "Срок", "Результат"], [
        ("1", "15.07.2026-22.07.2026", "Финализация ТЗ, модели данных, API-контрактов и матрицы ролей."),
        ("2", "23.07.2026-09.08.2026", "Завершение backend: content CRUD, аудит, ограничения запросов, тесты."),
        ("3", "10.08.2026-24.08.2026", "Завершение frontend: публичные ru/kk страницы, кабинет, admin UI, состояния ошибок и пустых данных."),
        ("4", "25.08.2026-07.09.2026", "Интеграционные испытания, миграции, исправление дефектов."),
        ("5", "08.09.2026-20.09.2026", "Security hardening, Docker/infra readiness, документация, backup/restore dry run."),
        ("6", "21.09.2026-30.09.2026", "Приемка, release candidate, передача продуктивной версии."),
    ], [800, 2200, 6360])
    h(doc, "2.5. Финансирование", 2)
    p(doc, "Финансирование работ выполняется из утвержденного бюджета проекта цифрового развития интернет-ресурса КНБ РК. Расчетная структура оплаты: 20% после утверждения ТЗ, 35% после готовности backend и миграций, 25% после готовности frontend и admin UI, 20% после приемки release candidate и передачи эксплуатационной документации.")
    h(doc, "2.6. Сокращения", 2)
    matrix(doc, ["Сокращение", "Значение"], [
        ("АС", "Автоматизированная система."),
        ("API", "Программный интерфейс системы."),
        ("БД", "База данных."),
        ("КНБ РК", "Комитет национальной безопасности Республики Казахстан."),
        ("ПДн", "Персональные данные."),
        ("RBAC", "Ролевая модель управления доступом."),
        ("JWT", "Токен доступа JSON Web Token."),
        ("ТЗ", "Техническое задание."),
        ("CRUD", "Создание, чтение, изменение и удаление записей."),
    ], [1900, 7460])
    doc.add_page_break()

    section(doc, 3, "Характеристика объектов автоматизации")
    h(doc, "3.1. Описание автоматизируемого процесса", 2)
    p(doc, "Автоматизируется официальный публичный контур взаимодействия Комитета с населением: публикация проверенных сведений о деятельности Комитета, пресс-сообщений, направлений работы, контактов и нормативных материалов; прием первичных заявок по вопросам работы и обучения; регистрация кандидатов; прохождение первичного психологического тестирования; служебная обработка обращений, кандидатских заявок, результатов тестов и материалов сайта.")
    h(doc, "3.2. Функции практики управления", 2)
    bullets(doc, [
        "Двуязычное официальное освещение деятельности Комитета на русском и казахском языках.",
        "Публикация разделов: о Комитете, деятельность, пресс-центр, документы, контакты, прием на службу, работа с населением, социальное обеспечение, противодействие терроризму, киберщит, защита государственных секретов, защита и охрана Государственной границы.",
        "Предоставление контактных данных: адрес центрального аппарата, канцелярия, прием граждан, телефоны доверия, пресс-служба, региональные контакты.",
        "Подача первичного обращения по вопросам работы или обучения с выдачей кода отслеживания.",
        "Создание учетной записи кандидата через Telegram с подтверждением телефона или через email/password.",
        "Административная обработка обращений, кандидатских заявок, результатов тестов и материалов портала.",
    ])
    h(doc, "3.3. Объект как техническая система", 2)
    p(doc, "Портал строится как web/API/DB-система. Frontend реализуется на Next.js App Router, React, TypeScript и Tailwind CSS. Backend реализуется на FastAPI, SQLAlchemy 2, Pydantic и Alembic. Основная БД - PostgreSQL 16. SQLite используется только локально разработчиками. Продуктивный запуск выполняется контейнерно: web, api, postgres, reverse proxy/ingress, мониторинг, логирование и резервное копирование.")
    h(doc, "3.4. Существующая практика контроля", 2)
    p(doc, "Исходное состояние проекта включает подключенные backend-сценарии авторизации, обращений, кабинета кандидата, admin panel и психологических тестов. Итоговая версия закрывает оставшиеся пробелы: content CRUD, полный аудит административных действий, системные тесты, production checks, security headers, request size limits, backup/restore, healthchecks и документацию.")
    h(doc, "3.5. Используемые средства", 2)
    kv(doc, [
        ("Frontend", "Next.js App Router, React, TypeScript, Tailwind CSS, lucide-react, Leaflet/react-leaflet."),
        ("Backend", "FastAPI, SQLAlchemy 2, Pydantic, Alembic, JWT, HttpOnly refresh cookie."),
        ("Database", "PostgreSQL 16 в Docker/production; SQLite только в локальной разработке через apps/api/.env."),
        ("Authentication", "Telegram login с подтверждением телефона через контакт бота; email/password как вторичный способ."),
        ("Admin access", "Двухступенчатый доступ: обычная portal session разрешенного пользователя и отдельный admin_session JWT."),
    ])
    h(doc, "3.6. Проблемы текущего состояния и закрывающие решения", 2)
    matrix(doc, ["Проблема", "Решение в настоящем ТЗ"], [
        ("Статические данные в frontend", "Перенести активные разделы в backend-backed content CRUD с локалями ru/kk."),
        ("Неполный аудит", "Логировать все create/update/delete/status/login-security действия."),
        ("Неполная security readiness", "Включить CORS allowlist, trusted hosts, CSRF для cookie flows, body limits, headers, startup secret checks."),
        ("Неполное тестирование", "Добавить backend tests и frontend typecheck/build/smoke для ru/kk flows."),
        ("Неполная эксплуатация", "Добавить healthchecks, non-root containers, env docs, backup/restore, logging и monitoring checklist."),
    ], [3200, 6160])
    h(doc, "3.7. Условия эксплуатации", 2)
    p(doc, "Продуктивная среда размещается в инфраструктуре с TLS 1.2+, WAF/rate limiting на периметре, закрытой сетью БД, отдельными staging и production окружениями, централизованным сбором логов, мониторингом доступности, ежедневными резервными копиями PostgreSQL и запретом публичного доступа к порту БД.")
    h(doc, "3.8. Предлагаемая структура обеспечения", 2)
    p(doc, "Структура обеспечения включает исходный код, миграции Alembic, CI checks, контейнеры web/api/db, инструкции запуска, инструкции эксплуатации, матрицу ролей, документацию API, журнал аудита, журнал попыток входа и эксплуатационный регламент.")
    doc.add_page_break()

    section(doc, 4, "Требования к системе в целом")
    p(doc, "Портал является официальным государственным интернет-ресурсом с двумя локалями: /ru и /kk. Все публичные материалы, формы, статусы и ошибки выводятся на языке выбранной локали. Английская локаль, отдельный /mobile сайт, поиск, загрузка файлов, eGov/ЭЦП/SMS/external identity integrations и managed document storage исключены из первой продуктивной версии.")
    h(doc, "4.1. Архитектурные требования", 2)
    bullets(doc, [
        "Frontend и API являются stateless-компонентами и масштабируются горизонтально.",
        "Access token хранится в browser session storage; refresh token хранится только в HttpOnly cookie.",
        "Refresh sessions хранятся в таблице RefreshSession, поддерживают rotation, revocation, срок жизни 30 дней и привязку к user_agent/ip для журнала.",
        "Все изменения схемы БД выполняются только через Alembic migrations; Base.metadata.create_all() в runtime запрещен.",
        "Все публичные и административные payload проходят Pydantic validation.",
        "Все запросы к БД выполняются через ORM/query parameterization без ручной сборки SQL из пользовательского ввода.",
    ])
    h(doc, "4.2. Функциональный состав первой версии", 2)
    matrix(doc, ["Модуль", "Состав"], [
        ("Public Portal", "Главная, о Комитете, направления, пресс-центр, документы, контакты, прием на службу, работа с населением, нормативная база."),
        ("Appeals", "Подача первичного обращения, tracking code, проверка статуса."),
        ("Candidate Account", "Регистрация, login, logout, refresh, профиль, заявка, статус, результаты тестов."),
        ("Psychological Testing", "Intro page, authenticated full-screen runner, timer, sections, answer persistence, results list."),
        ("Admin Panel", "Dashboard, appeals, candidate applications, psychological results, content CRUD, audit log."),
        ("Security", "RBAC, admin_session, throttling, CORS, trusted hosts, security headers, CSRF, request limits."),
    ], [1800, 7560])
    h(doc, "4.3. Нефункциональные параметры", 2)
    matrix(doc, ["Параметр", "Требование"], [
        ("Availability", "99,5% календарного месяца для web/API при исправной инфраструктуре."),
        ("API latency", "p95 до 1000 мс для чтения и до 2000 мс для записи при штатной нагрузке до 100 одновременных пользователей."),
        ("Backup", "Ежедневно в 02:00 Asia/Almaty; хранение ежедневных копий 14 дней, еженедельных 8 недель, ежемесячных 12 месяцев."),
        ("Recovery", "RPO 24 часа, RTO 4 часа."),
        ("Logging", "Application logs и audit logs хранятся 12 месяцев; security events - 24 месяца."),
        ("Accessibility", "Контрастность, keyboard navigation, readable focus states, semantic headings, формы с labels."),
    ], [2300, 7060])
    doc.add_page_break()

    section(doc, 5, "Требования к структуре и функционированию системы")
    h(doc, "5.1. Подсистемы", 2)
    matrix(doc, ["Подсистема", "Функции", "Данные"], [
        ("Публичный сайт", "Освещение деятельности Комитета, новости, направления, контакты, документы.", "Page, News, RegionOffice, static legal links."),
        ("Авторизация", "Telegram start/webhook/complete, email/password register/login, refresh, logout, /auth/me.", "User, TelegramLoginChallenge, RefreshSession, LoginAttempt."),
        ("Кандидат", "Кабинет, кандидатская заявка, статус, психологические результаты.", "CandidateApplication, PsychologicalTestResult."),
        ("Обращения", "Подача, tracking, статусы, модерация.", "Appeal, AuditLog."),
        ("Администрирование", "Dashboard, модерация, content CRUD, аудит.", "User, AuditLog, Page, News, RegionOffice."),
        ("Эксплуатация", "Health, migrations, logs, backup, monitoring.", "PostgreSQL backups, application logs, metrics."),
    ], [1500, 3600, 4260])
    h(doc, "5.2. Режимы функционирования", 2)
    bullets(doc, [
        "Штатный режим: web/API/DB доступны, пользователи проходят публичные и кабинетные сценарии.",
        "Режим обслуживания: публичный сайт показывает страницу технических работ; admin/API доступны только техническому оператору.",
        "Аварийный режим: запись новых обращений временно блокируется, чтение публичного контента сохраняется из кеша/reverse proxy, инцидент фиксируется в журнале.",
        "Staging режим: тестовая среда с отдельной БД и отдельными секретами; production данные не копируются без обезличивания.",
    ])
    h(doc, "5.3. Надежность", 2)
    bullets(doc, [
        "Все write-операции обращений, заявок, результатов тестов и admin actions выполняются транзакционно.",
        "Tracking code обращения и кандидатской заявки уникален и индексируется.",
        "Refresh-token rotation не допускает повторное использование старого refresh token.",
        "При недоступности API frontend показывает локализованную ошибку и кнопку повторной попытки.",
        "Миграции проверяются на пустой БД и на копии staging БД перед production deployment.",
    ])
    h(doc, "5.4. Безопасность", 2)
    bullets(doc, [
        "Пароли хэшируются стойким алгоритмом; plaintext password не логируется и не сохраняется.",
        "Telegram phone confirmation считается действительной только при contact.user_id == message.from.id.",
        "Webhook Telegram принимает запросы только с валидным secret token в production.",
        "Admin API требует admin_session claim; обычной роли пользователя недостаточно.",
        "CORS_ORIGINS и ALLOWED_HOSTS задаются allowlist; wildcard в production запрещен.",
        "Cookie refresh устанавливается с HttpOnly, Secure и SameSite=Lax в production.",
        "CSRF token применяется к cookie-backed state-changing flows.",
        "Request body limit: 256 KB для публичных форм, 1 MB для admin content API.",
        "Security headers: Content-Security-Policy, X-Frame-Options/Frame-Ancestors, X-Content-Type-Options, Referrer-Policy, Permissions-Policy, HSTS на TLS ingress.",
    ])
    h(doc, "5.5. Персонал и режим работы", 2)
    matrix(doc, ["Роль", "Права", "Режим"], [
        ("Admin", "Полный доступ к admin panel, content CRUD, users status, audit view, moderation.", "08:00-20:00 рабочие дни; emergency доступ 24/7."),
        ("Moderator", "Обработка обращений, кандидатских заявок, просмотр результатов тестов без удаления данных.", "08:00-20:00 рабочие дни."),
        ("Candidate", "Собственный кабинет, собственные результаты, подача/просмотр собственных данных.", "24/7."),
        ("Technical operator", "Deployment, logs, backup/restore, monitoring; без редактирования смыслового контента.", "24/7 для инцидентов."),
    ], [1600, 5400, 2360])
    h(doc, "5.6. Эргономика", 2)
    p(doc, "Интерфейс использует сдержанную государственную визуальную систему, читаемую типографику, ясные labels, видимые focus states, понятные статусы, адаптивную компоновку для mobile и desktop без отдельных мобильных маршрутов. Формы имеют inline validation, summary ошибок и подтверждение успешной отправки.")
    h(doc, "5.7. Эксплуатация", 2)
    p(doc, "Контейнеры web и api запускаются non-root пользователями. Healthchecks: /health для API и HTTP 200 для web. PostgreSQL не публикуется наружу в production. Все секреты поступают из environment/secret manager. Docker images собираются воспроизводимо и маркируются commit SHA.")
    h(doc, "5.8. Сохранность информации", 2)
    p(doc, "Потеря подтвержденной записи обращения, заявки, результата теста, refresh session или audit log является критическим дефектом. Резервное копирование PostgreSQL выполняется ежедневно; восстановление проверяется ежемесячно на isolated restore environment.")
    h(doc, "5.9. Патентная чистота и лицензии", 2)
    p(doc, "Разрешены зависимости с лицензиями MIT, Apache-2.0, BSD, ISC и совместимыми permissive лицензиями. GPL/AGPL зависимости в runtime запрещены без отдельного письменного разрешения. Список зависимостей фиксируется в package-lock/requirements и проверяется перед релизом.")
    doc.add_page_break()

    section(doc, 6, "Требования к функциям (задачам)")
    h(doc, "6.1. Публичные функции", 2)
    bullets(doc, [
        "Главная страница выводит краткое назначение Комитета, пресс-материалы, направления деятельности, быстрые ссылки на контакты, прием на службу и обращения.",
        "Раздел «О Комитете» описывает руководство единой системой органов национальной безопасности, разведывательную, контрразведывательную, оперативно-розыскную деятельность, охрану Государственной границы, правительственную связь, специальное назначение и защиту государственных секретов.",
        "Раздел «Пресс-центр» содержит новости и официальные сообщения с датой публикации, локалями ru/kk, статусом draft/published/archived и slug.",
        "Раздел «Контакты» содержит адрес, телефоны канцелярии, приема граждан, доверия, пресс-службы, карту и региональные офисы.",
        "Раздел «Документы/Нормативная база» содержит статические ссылки и справочную информацию без загрузки файлов.",
        "Раздел «Прием на службу» содержит условия, этапы, список первичных сведений и переход к регистрации кандидата.",
    ])
    h(doc, "6.2. Авторизация и кабинет", 2)
    numbers(doc, [
        "Пользователь выбирает вход через Telegram или email/password.",
        "Telegram flow создает challenge, открывает bot deep link, принимает webhook, проверяет phone contact ownership, завершает вход и создает User(role=candidate).",
        "Email/password registration собирает first name, last name, birth date, phone, email, password и создает User(role=candidate) + CandidateApplication(status=submitted).",
        "После входа frontend получает access token, сохраняет его в session storage и восстанавливает сессию через /auth/refresh при истечении access token.",
        "Logout отзывает refresh session, очищает cookie и очищает session storage.",
    ])
    h(doc, "6.3. Обращения", 2)
    bullets(doc, [
        "Публичная форма принимает ФИО, email, телефон, тему, сообщение и необязательный ИИН.",
        "После отправки система создает Appeal(status=received), генерирует tracking code и показывает инструкцию: документы направляются на hr@knb.gov.kz.",
        "Статусы обращения: received, in_review, answered, rejected.",
        "Статус lookup доступен по tracking code без раскрытия персональных данных сверх минимума.",
    ])
    h(doc, "6.4. Психологическое тестирование", 2)
    bullets(doc, [
        "Intro route: /[locale]/psychological-testing.",
        "Runner route: /[locale]/psychological-testing/primary-selection.",
        "Доступ только после успешного /auth/me.",
        "Тест запускается в full-screen mode без обычного header/footer портала.",
        "Таймер: 60 минут активного времени, пауза на экранах инструкций.",
        "Пагинация: 10 вопросов на страницу.",
        "Сохраняются answers, sections, answered_questions, duration_seconds, remaining_seconds, submitted_at.",
        "Автоматический scoring не выполняется до включения утвержденных ключей в отдельной версии.",
    ])
    h(doc, "6.5. Администрирование", 2)
    bullets(doc, [
        "Root /admin сначала требует обычную авторизацию пользователя с email из ADMIN_PORTAL_ALLOWED_USER_EMAIL.",
        "Второй шаг требует ADMIN_PANEL_EMAIL и ADMIN_PANEL_PASSWORD_HASH и выдает admin_session JWT.",
        "Dashboard показывает счетчики обращений, кандидатских заявок, результатов тестов, опубликованных страниц и последних действий.",
        "Appeals view поддерживает список, фильтр по статусу, просмотр деталей, назначение статуса и moderator comment.",
        "Candidate applications view поддерживает список, фильтр по статусу, просмотр карточки, смену статуса и комментарий.",
        "Testing tab показывает результаты PsychologicalTestResult с кандидатом, датой, длительностью и секционными summary.",
        "Content CRUD управляет Page, News и RegionOffice с draft/published/archived состояниями.",
        "Audit tab показывает actor, action, entity, entity_id, ip_address и created_at.",
    ])
    h(doc, "6.6. Критерии отказов", 2)
    matrix(doc, ["Функция", "Отказ", "Блокировка приемки"], [
        ("Auth", "Вход без подтвержденной сессии или обход admin_session.", "Да."),
        ("Appeal", "Нет tracking code после успешной отправки.", "Да."),
        ("Candidate", "Кандидат видит чужие данные.", "Да."),
        ("Test", "Результат завершенного теста не сохраняется.", "Да."),
        ("Content", "Опубликованный ru/kk контент отображается не в той локали.", "Да."),
        ("Audit", "Admin status/content action не попал в AuditLog.", "Да."),
    ], [1400, 4400, 3560])
    doc.add_page_break()

    section(doc, 7, "Требования к видам обеспечения")
    h(doc, "7.1. Математическое обеспечение", 2)
    p(doc, "Математическое обеспечение первой версии ограничивается расчетом таймера, количества отвеченных вопросов, длительности прохождения, агрегированных секционных summary и счетчиков dashboard. Психологическая интерпретация не рассчитывается: система сохраняет первичные ответы и метаданные для последующей экспертной оценки.")
    h(doc, "7.2. Информационное обеспечение", 2)
    matrix(doc, ["Сущность", "Назначение", "Ключевые поля"], [
        ("User", "Учетная запись.", "email, full_name, telegram_id, phone, role, is_active, is_blocked."),
        ("CandidateApplication", "Кандидатская заявка.", "tracking_code, ФИО, birth_date, phone, region, education_level, desired_direction, status."),
        ("Appeal", "Первичное обращение.", "tracking_code, full_name, email, phone, subject, message, status."),
        ("Page/News", "Управляемый контент.", "locale fields, slug, title, body, status, published_at."),
        ("RegionOffice", "Контакты регионов.", "service, name_ru/kk, region_ru/kk, phones, latitude, longitude."),
        ("AuditLog", "Административный журнал.", "actor_id, action, entity, entity_id, ip_address, created_at."),
    ], [1900, 2700, 4760])
    h(doc, "7.3. Лингвистическое обеспечение", 2)
    p(doc, "Поддерживаются только русский и казахский языки. Все public routes имеют префикс /ru или /kk. Переключение языка сохраняет эквивалентный маршрут. Контентные сущности хранят русскую и казахскую версии; публикация разрешена только при заполнении обеих версий.")
    h(doc, "7.4. Программное обеспечение", 2)
    bullets(doc, [
        "Node.js 20+ для frontend build/runtime.",
        "Python 3.12+ для backend.",
        "PostgreSQL 16 для production.",
        "FastAPI app.main:app через uvicorn/gunicorn или эквивалентный ASGI runtime.",
        "Next.js build/start для web.",
        "Alembic upgrade head перед запуском новой версии API.",
    ])
    h(doc, "7.5. Техническое обеспечение", 2)
    bullets(doc, [
        "2 vCPU / 4 GB RAM для web на стартовой нагрузке.",
        "2 vCPU / 4 GB RAM для API на стартовой нагрузке.",
        "2 vCPU / 8 GB RAM / SSD для PostgreSQL.",
        "Отдельный persistent volume для PostgreSQL.",
        "Отдельное хранилище резервных копий вне основного DB volume.",
    ])
    h(doc, "7.6. Метрологическое обеспечение", 2)
    p(doc, "Метрологические средства не используются. Контролируемые показатели: uptime, API p95 latency, error rate, DB disk usage, backup success, restore duration, failed login count, blocked requests, audit event count.")
    h(doc, "7.7. Организационное и методическое обеспечение", 2)
    bullets(doc, [
        "Регламент публикации официальных материалов.",
        "Регламент обработки первичных обращений и кандидатских заявок.",
        "Регламент работы с персональными данными.",
        "Регламент реагирования на инциденты.",
        "Инструкция администратора, модератора и технического оператора.",
    ])
    doc.add_page_break()

    section(doc, 8, "Состав и содержание работ по созданию системы")
    h(doc, "8.1. Стадии", 2)
    matrix(doc, ["Стадия", "Работы", "Артефакты"], [
        ("ТЗ", "Финализация требований, источников, scope и приемки.", "ТЗ v2.0, матрица функций, список источников."),
        ("Проектирование", "Data model, API, RBAC, UX flows, deployment plan.", "ERD, API contract, route map, security checklist."),
        ("Разработка backend", "Модели, миграции, endpoints, validation, audit, tests.", "FastAPI code, Alembic migrations, pytest suite."),
        ("Разработка frontend", "ru/kk pages, forms, cabinet, admin panel, states.", "Next.js app, typed helpers, components."),
        ("Security hardening", "Headers, CORS, CSRF, rate limits, body limits, startup checks.", "Config, tests, hardening notes."),
        ("Эксплуатация", "Docker, healthchecks, logs, backup/restore, env docs.", "Compose/manifests, runbook, .env.example."),
        ("Приемка", "Functional tests, smoke tests, defect closure, release.", "Test report, release notes, acceptance act."),
    ], [1400, 4700, 3260])
    h(doc, "8.2. Документы на выходе", 2)
    bullets(doc, [
        "Техническое задание v2.0.",
        "Архитектурное описание.",
        "Описание API.",
        "Описание БД и миграций.",
        "Инструкция администратора.",
        "Инструкция модератора.",
        "Руководство развертывания.",
        "Руководство backup/restore.",
        "Программа и методика испытаний.",
        "Протокол испытаний.",
        "Release notes.",
    ])
    h(doc, "8.3. Экспертиза документации", 2)
    p(doc, "Экспертиза выполняется проектной комиссией из представителя заказчика, разработчика, технического оператора и специалиста информационной безопасности. Комиссия проверяет полноту функций, соответствие scope, защищенность, качество документации, воспроизводимость deployment и отсутствие незакрытых критических дефектов.")
    h(doc, "8.4. Работы по надежности", 2)
    bullets(doc, [
        "Unit tests backend для auth, sessions, RBAC, appeals, candidates, admin actions.",
        "Integration tests API на PostgreSQL.",
        "Frontend typecheck и production build.",
        "Smoke tests ru/kk main flows.",
        "Monthly restore drill.",
        "Dependency audit перед релизом.",
    ])
    h(doc, "8.5. Работы по метрологическому обеспечению", 2)
    p(doc, "Вместо метрологических работ выполняется настройка измеримых эксплуатационных метрик и алертов: uptime, latency, error rate, disk usage, backup success, failed login spikes, rate-limit triggers.")
    doc.add_page_break()

    section(doc, 9, "Порядок контроля и приемки системы")
    h(doc, "9.1. Верификация", 2)
    bullets(doc, [
        "Все 12 разделов ТЗ закрыты конкретными требованиями.",
        "Нет out-of-scope интеграций: eGov, ЭЦП, SMS, внешние identity providers, file upload, search, English locale, /mobile.",
        "Все public ru/kk маршруты открываются и показывают локализованный контент.",
        "Auth, refresh, logout и /auth/me работают end-to-end.",
        "Appeal create/status lookup работает end-to-end.",
        "Candidate dashboard показывает собственные данные и результаты.",
        "Admin panel требует отдельный admin_session.",
        "Все admin status/content changes попадают в AuditLog.",
        "Alembic upgrade head выполняется на чистой БД.",
        "Frontend typecheck и build завершаются успешно.",
    ])
    h(doc, "9.2. Аттестация", 2)
    p(doc, "Система считается аттестованной для опытно-промышленной эксплуатации после успешного прохождения приемочного сценария, закрытия всех критических и высоких дефектов, готовности backup/restore и подтверждения security baseline.")
    h(doc, "9.3. Совместная проверка", 2)
    matrix(doc, ["Проверка", "Критерий успеха"], [
        ("Публичный портал", "Главная, пресс-центр, документы, контакты и прием на службу доступны на /ru и /kk."),
        ("Обращение", "Создается Appeal, отображается tracking code, lookup возвращает статус."),
        ("Кандидат", "Регистрация создает User и CandidateApplication, кабинет доступен после входа."),
        ("Тест", "Runner сохраняет PsychologicalTestResult и показывает его в кабинете/admin."),
        ("Админ", "Admin/moderator может менять статусы в рамках прав; действие записано в AuditLog."),
        ("Эксплуатация", "Healthchecks зеленые, backup создан, restore проверен."),
    ], [2300, 7060])
    h(doc, "9.4. Аудит", 2)
    p(doc, "Аудит приемки проверяет соответствие репозитория ТЗ, отсутствие hardcoded production secrets, отсутствие demo auth, наличие миграций, журналирование admin actions, корректность зависимостей и reproducible deployment.")
    h(doc, "9.5. Разрешение проблем", 2)
    p(doc, "Дефекты классифицируются как Critical, High, Medium, Low. Critical и High блокируют релиз. Medium допускаются только при наличии зафиксированного workaround и срока закрытия до 14 календарных дней. Low включаются в backlog гарантийного периода.")
    h(doc, "9.6. Приемка по стадиям", 2)
    p(doc, "Приемка каждой стадии оформляется протоколом. Итоговая приемка оформляется актом после демонстрации всех приемочных сценариев, передачи документов, release tag, инструкций эксплуатации и подтверждения backup/restore.")
    doc.add_page_break()

    section(doc, 10, "Подготовка объекта автоматизации к вводу системы в действие")
    h(doc, "10.1. Подготовка информации", 2)
    bullets(doc, [
        "Подготовить двуязычные тексты разделов: о Комитете, деятельность, пресс-центр, документы, контакты, прием на службу, работа с населением.",
        "Заполнить RegionOffice для центрального аппарата и региональных контактов.",
        "Проверить все телефонные номера, email и адреса по официальному ресурсу GOV.KZ.",
        "Очистить seed/demo данные перед production.",
        "Проверить, что published контент не содержит черновиков и внутренних комментариев.",
    ])
    h(doc, "10.2. Организационные изменения", 2)
    bullets(doc, [
        "Назначить двух администраторов, четырех модераторов, одного технического оператора и одного ответственного за публикации.",
        "Утвердить порядок смены статусов обращений и кандидатских заявок.",
        "Утвердить регламент публикации материалов пресс-службы.",
        "Утвердить порядок блокировки пользователя и отзыва refresh sessions.",
        "Утвердить порядок хранения и удаления персональных данных.",
    ])
    h(doc, "10.3. Условия функционирования", 2)
    p(doc, "До запуска создаются production домен, TLS certificate, CORS allowlist, ALLOWED_HOSTS, JWT_SECRET, REFRESH_COOKIE_SECRET, TELEGRAM_BOT_TOKEN, TELEGRAM_WEBHOOK_SECRET, ADMIN_PORTAL_ALLOWED_USER_EMAIL, ADMIN_PANEL_EMAIL, ADMIN_PANEL_PASSWORD_HASH, DATABASE_URL, backup credentials и monitoring endpoints. Все значения размещаются в secret manager или защищенных environment variables.")
    h(doc, "10.4. Службы эксплуатации", 2)
    p(doc, "Эксплуатационная зона ответственности разделяется на content owner, moderator team, technical operator и security reviewer. Каждый инцидент получает номер, severity, владельца, срок реакции и итоговое решение.")
    h(doc, "10.5. Обучение", 2)
    p(doc, "Администраторы и модераторы проходят 2-часовое обучение: вход, второй admin login, списки, фильтры, смена статусов, комментарии, content CRUD, audit log, ошибки, правила ПДн. Технический оператор проходит 3-часовое обучение: deployment, migrations, logs, healthchecks, backup, restore, rollback.")
    doc.add_page_break()

    section(doc, 11, "Требования к документированию")
    h(doc, "11.1. Комплект документов", 2)
    bullets(doc, [
        "ТЗ v2.0 в DOCX и PDF.",
        "README с локальным, Docker и production-подобным запуском.",
        ".env.example для web и API без секретов.",
        "ARCHITECTURE.md с актуальной схемой.",
        "API reference с endpoints, request/response и auth requirements.",
        "Database schema reference с миграциями.",
        "Admin/moderator manual.",
        "Deployment and operations runbook.",
        "Backup/restore runbook.",
        "Security hardening checklist.",
        "Test plan and acceptance report.",
    ])
    h(doc, "11.2. Машинные носители", 2)
    p(doc, "Исходный код, документация, миграции, тесты, Docker-конфигурация и release notes передаются в Git-репозитории. Итоговые документы передаются в DOCX/PDF/Markdown. Production-секреты в репозиторий не помещаются.")
    h(doc, "11.3. Документирование компонентов", 2)
    p(doc, "Каждый runtime component имеет описание назначения, команд запуска, переменных окружения, healthcheck, логов, ошибок и recovery procedure. Каждая интеграционная точка описана в API reference.")
    h(doc, "11.4. Дополнительные требования", 2)
    bullets(doc, [
        "Документы пишутся на русском языке; пользовательские публичные тексты также ведутся на казахском.",
        "Все screenshots в инструкциях скрывают персональные данные.",
        "Каждая инструкция содержит дату версии и связанный release tag.",
        "Документация обновляется в том же pull request, где меняется соответствующее поведение.",
    ])
    doc.add_page_break()

    section(doc, 12, "Источники разработки")
    h(doc, "12.1. Открытые и проектные источники", 2)
    bullets(doc, [
        "Официальный ресурс КНБ РК на GOV.KZ: https://www.gov.kz/memleket/entities/knb?lang=ru.",
        "Казахская версия официального ресурса КНБ РК на GOV.KZ: https://www.gov.kz/memleket/entities/knb?lang=kk.",
        "Разделы официального ресурса: пресс-центр, направления, документы, контакты, прием на службу, работа с населением, социальные медиа и полезные ссылки.",
        "Открытые ссылки GOV.KZ на e-Otinish, eGov и eLicense использованы как пример внешних справочных переходов без API-интеграции.",
        "Проектный репозиторий: apps/web, apps/api, docs/ARCHITECTURE.md, SECURITY_ROADMAP.md, docker-compose.yml, Alembic migrations.",
        "Исходный шаблон: техническое_задание_2026-07-13.docx.",
        "ГОСТ 34.601 как структура стадий создания автоматизированной системы.",
    ])
    h(doc, "12.2. Зафиксированные проектные решения", 2)
    bullets(doc, [
        "Две локали: ru и kk.",
        "PostgreSQL в production, SQLite только локально.",
        "Telegram и email/password - единственные способы auth первой версии.",
        "Admin-panel login отделен от ordinary auth.",
        "Подача файлов, ЭЦП, SMS, eGov API, search и English locale не входят в версию 1.0.",
        "Content CRUD реализуется только для активных публичных разделов: Page, News, RegionOffice.",
        "Психологические результаты сохраняются без автоматического scoring.",
    ])
    p(doc, "Настоящее ТЗ является закрытой полной редакцией требований для реализации первой продуктивной версии портала КНБ РК. Все спорные вопросы scope в этой версии решены текстом настоящего документа.", bold=True)

    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
