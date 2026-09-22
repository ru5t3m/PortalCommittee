from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("Техническое_задание_портал_КНБ_2026-07-13.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(89, 89, 89)
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "F2F4F7"


def set_font(run, name="Calibri", size=11, color=BLACK, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def paragraph_border_bottom(paragraph, color="2E74B5", size="12", space="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tbl_cell_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def add_para(doc, text="", style=None, bold=False, italic=False, color=BLACK, size=11, align=None, before=0, after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        set_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullets(doc, items):
    for item in items:
        p = add_para(doc, style="List Bullet", after=4)
        run = p.add_run(item)
        set_font(run, size=11)


def add_numbers(doc, items):
    for item in items:
        p = add_para(doc, style="List Number", after=4)
        run = p.add_run(item)
        set_font(run, size=11)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt({1: 16, 2: 12, 3: 8}[level])
    p.paragraph_format.space_after = Pt({1: 8, 2: 6, 3: 4}[level])
    run = p.add_run(text)
    set_font(run, size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return p


def add_kv_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    set_cell_margins(table)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        set_cell_shading(cells[0], HEADER_FILL)
        for cell in cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    set_font(r, size=10.5, bold=(cell is cells[0]))
    set_table_width(table, [2200, 7160])
    add_para(doc, "")
    return table


def add_matrix_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_cell_margins(table)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, HEADER_FILL)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_font(r, size=10, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i != 0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    set_font(r, size=9.5)
    set_table_width(table, widths)
    add_para(doc, "")
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.bold = True
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = BLUE
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.color.rgb = BLUE
    styles["Heading 3"].font.size = Pt(12)
    styles["Heading 3"].font.color.rgb = DARK_BLUE

    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    r = hp.add_run("Техническое задание | Официальный интернет-портал КНБ РК")
    set_font(r, size=9, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Конфиденциальность: для проектного согласования")
    set_font(fr, size=9, color=MUTED)


def add_title_page(doc):
    add_para(doc, "Техническое задание", size=24, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_para(
        doc,
        "на создание и развитие автоматизированной информационной системы официального интернет-портала Комитета национальной безопасности Республики Казахстан",
        size=15,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=18,
    )
    add_para(doc, "Раздел 1 из 12", size=11, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=20)
    add_kv_table(
        doc,
        [
            ("Объект автоматизации", "Официальный интернет-портал КНБ РК с публичным сайтом, кабинетом кандидата и административной панелью."),
            ("Заказчик", "Комитет национальной безопасности Республики Казахстан либо уполномоченное структурное подразделение Заказчика."),
            ("Разработчик", "Исполнитель определяется договором на разработку, внедрение и сопровождение системы."),
            ("Версия документа", "1.0 от 13.07.2026."),
            ("Статус", "Проект для согласования требований, планирования работ и последующей приемки."),
        ],
    )
    add_para(
        doc,
        "Документ подготовлен по структуре шаблона технического задания и описывает функциональные, технические, организационные и приемочные требования к порталу как к продуктивному интернет-ресурсу государственного органа.",
        after=10,
    )
    p = add_para(doc, "")
    paragraph_border_bottom(p)
    add_heading(doc, "Аннотация", 1)
    add_para(
        doc,
        "Настоящее техническое задание определяет назначение, цели, состав функций, требования к надежности, безопасности, видам обеспечения, порядку разработки, документирования, контроля и приемки автоматизированной информационной системы официального интернет-портала КНБ РК.",
    )
    add_para(
        doc,
        "Система предназначена для официального информирования граждан, приема первичных обращений по вопросам работы и обучения, регистрации кандидатов, проведения первичного психологического тестирования, ведения кабинета кандидата, а также для модерирования обращений, кандидатских заявок и материалов портала сотрудниками с ролями Администратор и Модератор.",
    )
    add_para(
        doc,
        "Проект не предусматривает интеграции с eGov, ЭЦП, SMS, внешними поставщиками идентификации, загрузки файлов, управляемого хранилища документов, поиска, английской локали и отдельной мобильной версии по маршрутам вида /mobile.",
        bold=True,
    )
    doc.add_page_break()


def section_title(doc, number, title, subtitle=None):
    add_para(doc, f"Раздел {number} из 12", size=10.5, color=MUTED, after=3)
    add_heading(doc, title, 1)
    if subtitle:
        add_para(doc, subtitle, italic=True, color=MUTED, after=8)


def build_doc():
    doc = Document()
    configure_document(doc)
    add_title_page(doc)

    section_title(doc, 2, "Назначение и цели создания системы")
    add_heading(doc, "2.1. Шифр темы или шифр (номер) договора", 2)
    add_para(doc, "Условное наименование темы: «Официальный интернет-портал КНБ РК». Шифр проекта для внутреннего учета: KNB-PORTAL-2026. Номер договора, реквизиты государственного задания и бюджетной программы указываются Заказчиком после утверждения закупочной или проектной документации.")
    add_heading(doc, "2.2. Наименование предприятий разработчика и заказчика", 2)
    add_kv_table(doc, [("Заказчик", "Комитет национальной безопасности Республики Казахстан или уполномоченное подразделение."), ("Пользователи Заказчика", "Администраторы, модераторы, редакторы и ответственные сотрудники кадрового/организационного направления."), ("Разработчик", "Организация-исполнитель, определяемая договором."), ("Публичные пользователи", "Граждане, кандидаты на работу или обучение, посетители официального портала.")])
    add_heading(doc, "2.3. Документы, на основании которых создается система", 2)
    add_bullets(doc, ["Договор, техническая спецификация, календарный план и иные приложения, утвержденные Заказчиком.", "Законодательство Республики Казахстан об информатизации, персональных данных, государственных интернет-ресурсах и информационной безопасности.", "Внутренние организационно-распорядительные документы Заказчика, регламентирующие публикацию информации, обработку обращений и работу с кандидатами.", "ГОСТ 34.601 и применимые практики жизненного цикла автоматизированных систем.", "Архитектурная документация проекта, roadmap безопасности и фактический исходный код репозитория."])
    add_heading(doc, "2.4. Плановые сроки начала и окончания работ", 2)
    add_para(doc, "Работы выполняются поэтапно. Конкретные календарные даты фиксируются в договоре и плане-графике. Базовый ориентир: начало работ после утверждения ТЗ; завершение первой продуктивной версии после прохождения функциональной приемки, проверки безопасности, устранения критических дефектов и подготовки эксплуатационной документации.")
    add_heading(doc, "2.5. Источники и порядок финансирования", 2)
    add_para(doc, "Финансирование осуществляется в порядке, определенном Заказчиком и договором. Порядок оплаты рекомендуется привязать к приемке стадий: проектирование, разработка, испытания, внедрение, гарантийное сопровождение.")
    add_heading(doc, "2.6. Список используемых сокращений", 2)
    add_matrix_table(doc, ["Сокращение", "Расшифровка"], [["АС", "Автоматизированная система."], ["API", "Программный интерфейс взаимодействия компонентов."], ["БД", "База данных."], ["JWT", "Токен доступа JSON Web Token."], ["RBAC", "Ролевая модель доступа."], ["КНБ РК", "Комитет национальной безопасности Республики Казахстан."], ["ПДн", "Персональные данные."], ["ТЗ", "Техническое задание."], ["UI/UX", "Пользовательский интерфейс и пользовательский опыт."]], [1800, 7560])
    doc.add_page_break()

    section_title(doc, 3, "Характеристика объектов автоматизации")
    add_heading(doc, "3.1. Краткое описание процесса", 2)
    add_para(doc, "Объектом автоматизации является процесс эксплуатации официального интернет-ресурса: публикация сведений для граждан, прием первичных заявок на работу и обучение, регистрация кандидатов, проведение первичного психологического тестирования, сопровождение статусов заявок и административная обработка поступающих данных.")
    add_heading(doc, "3.2. Задачи и функции, решаемые в практике управления", 2)
    add_bullets(doc, ["Официальное двуязычное информирование посетителей на русском и казахском языках.", "Прием первичных обращений с выдачей кода отслеживания.", "Создание и сопровождение учетной записи кандидата.", "Сбор минимально необходимых кандидатских данных без загрузки файлов.", "Отображение результатов психологического тестирования кандидату и администраторам.", "Модерация обращений и кандидатских заявок.", "Ведение аудита административных действий и журналов входа."])
    add_heading(doc, "3.3. Описание объекта управления", 2)
    add_para(doc, "Портал является веб-приложением с публичным frontend, backend API и реляционной базой данных. Статические и динамические характеристики объекта определяются количеством посетителей, числом одновременных сессий, объемом обращений, частотой изменения контента, требованиями к доступности и уровнем защиты персональных данных.")
    add_heading(doc, "3.4. Существующая практика контроля и управления", 2)
    add_para(doc, "До завершения автоматизации часть информационных материалов и процессов может поддерживаться статически или вручную. Существующая реализация уже содержит ряд подключенных backend-сценариев: регистрацию, авторизацию, обращения, администрирование, психологические тесты и аудит отдельных действий. Контент CRUD для активных разделов подлежит завершению.")
    add_heading(doc, "3.5. Технические, информационные и математические средства", 2)
    add_para(doc, "Используются Next.js App Router, React, TypeScript, Tailwind CSS, lucide-react, Leaflet/react-leaflet, FastAPI, SQLAlchemy 2, Pydantic, Alembic, PostgreSQL в Docker/production и SQLite только как локальный fallback. Для сессий применяются JWT access tokens и refresh-token sessions с HttpOnly cookie.")
    add_heading(doc, "3.6. Достоинства и недостатки существующей практики", 2)
    add_matrix_table(doc, ["Аспект", "Текущее состояние", "Требуемое развитие"], [["Функциональность", "Основные публичные и административные сценарии частично реализованы.", "Завершить content CRUD, состояния ошибок/пустых данных и покрытие тестами."], ["Безопасность", "Есть RBAC, refresh-сессии, Telegram phone ownership check, trusted hosts/CORS.", "Довести rate limiting, CSRF, security headers, request limits, prod secret checks и централизованные ошибки."], ["Эксплуатация", "Есть Dockerfiles и docker-compose.", "Добавить production env documentation, healthchecks, non-root containers, backup/restore и logging."]], [1900, 3730, 3730])
    add_heading(doc, "3.7. Условия эксплуатации объекта автоматизации", 2)
    add_para(doc, "Система эксплуатируется в защищенной серверной инфраструктуре Заказчика или уполномоченной hosting/reviewing party. Предполагается HTTPS на доверенном ingress/reverse proxy, WAF/rate limiting на периметре, сетевое разделение web/API/DB, закрытый доступ к БД из публичного интернета, централизованные логи, резервное копирование и мониторинг.")
    add_heading(doc, "3.8. Предлагаемая структура обеспечения", 2)
    add_para(doc, "Предлагаемая структура включает frontend-приложение, backend API, PostgreSQL, механизмы миграций, систему ролей, журналы аудита, документацию, тесты и эксплуатационные инструкции. Все изменения схемы данных выполняются через Alembic; создание таблиц при старте приложения запрещается.")
    doc.add_page_break()

    section_title(doc, 4, "Требования к системе в целом")
    add_para(doc, "Система должна быть продуктивным, проверяемым и сопровождаемым интернет-ресурсом государственного органа. Доступные локали: /ru и /kk. Поведение, маршруты, бизнес-правила и состав функций должны быть согласованы между локалями. Английская локаль и отдельные мобильные маршруты не создаются.")
    add_bullets(doc, ["Frontend должен обращаться к API через типизированные helper-функции и использовать переменную NEXT_PUBLIC_API_URL.", "Backend должен предоставлять REST API с Pydantic-валидацией входных данных.", "Роли ограничиваются Admin, Moderator и Candidate, если Заказчик письменно не расширит модель.", "Система не должна содержать demo-auth, hardcoded credentials, тестовые пароли или production-секреты в репозитории.", "Пользовательские сценарии должны иметь понятные loading, empty и error states.", "Все существенные административные операции должны журналироваться в AuditLog."])

    section_title(doc, 5, "Требования к структуре и функционированию системы")
    add_heading(doc, "5.1. Структура и функционирование", 2)
    add_matrix_table(doc, ["Подсистема", "Назначение", "Ключевые требования"], [["Публичный портал", "Информирование граждан и публикация официального контента.", "Две локали /ru и /kk, адаптивный UX, отсутствие поиска до изменения scope."], ["Кабинет кандидата", "Регистрация, профиль, статус заявки, результаты тестов.", "Telegram/email auth, восстановление сессии, раздельное хранение auth и candidate data."], ["Обращения", "Прием первичных work/study applications.", "Подача заявки, tracking code, lookup статуса, инструкция отправки документов на утвержденный email."], ["Психологическое тестирование", "Полноэкранный authenticated runner.", "Таймер, инструкции, пагинация, сохранение результатов, без автоматического scoring до ключей."], ["Админ-панель", "Модерация обращений, заявок, результатов и контента.", "Отдельный admin-session login, RBAC, аудит, пустые/ошибочные состояния."], ["API и БД", "Бизнес-логика и хранение данных.", "FastAPI, SQLAlchemy, Alembic, PostgreSQL, SQLite только локально."]], [1500, 3000, 4860])
    add_heading(doc, "5.2. Надежность", 2)
    add_bullets(doc, ["Доступность продуктивной версии определяется SLA договора; рекомендуемый целевой уровень для публичного портала - не ниже 99,5% в месяц при условии корректной инфраструктуры.", "Сессии refresh-token должны храниться серверно, поддерживать rotation/revocation и истечение срока действия.", "При отказе frontend должен показывать корректные сообщения и не раскрывать внутренние ошибки.", "Миграции БД должны быть обратимо планируемыми и предварительно проверенными на staging.", "Критические операции администрирования должны быть атомарными на уровне транзакций БД."])
    add_heading(doc, "5.3. Безопасность", 2)
    add_para(doc, "Требования безопасности включают strong authentication, RBAC, Telegram webhook secret validation, Telegram contact ownership check, brute-force/rate-limit protection, strict CORS allowlist, trusted hosts, security headers, CSRF для cookie-backed flows где применимо, request body size limits, production secret validation, централизованные безопасные ошибки, XSS prevention и ORM parameterization.")
    add_heading(doc, "5.4. Численность и квалификация персонала", 2)
    add_bullets(doc, ["Администратор: управление настройками, контентом, пользователями/ролями в пределах утвержденных функций, просмотр журналов.", "Модератор: обработка обращений, кандидатских заявок и материалов в пределах назначенных прав.", "Технический специалист: развертывание, резервное копирование, мониторинг, обновления и восстановление.", "Публичный кандидат: работа только со своей учетной записью, заявкой и результатами тестов."])
    add_heading(doc, "5.5. Показатели назначения", 2)
    add_bullets(doc, ["Подача обращения должна завершаться выдачей tracking code.", "Авторизованный кандидат должен видеть данные своего кабинета и результаты тестов.", "Администратор должен видеть актуальные списки обращений, заявок и результатов тестов.", "Портал должен корректно работать на основных desktop/mobile viewport без отдельных /mobile маршрутов.", "Система должна поддерживать дальнейшее расширение content CRUD без изменения публичных URL."])
    add_heading(doc, "5.6. Эргономика и техническая эстетика", 2)
    add_para(doc, "Интерфейс должен соответствовать официальному государственному ресурсу: сдержанная визуальная система, ясная типографика, понятные действия, крупные touch targets на мобильных устройствах, предсказуемая навигация, отсутствие рекламной стилистики и недопустимых декоративных элементов.")
    add_heading(doc, "5.7. Эксплуатация, обслуживание и хранение компонентов", 2)
    add_para(doc, "Компоненты web, API и DB разворачиваются контейнерно или эквивалентно. Должны быть описаны переменные окружения, порядок запуска миграций, healthchecks, backup/restore PostgreSQL, логирование, ротация логов и порядок установки обновлений.")
    add_heading(doc, "5.8. Защита информации от несанкционированного доступа", 2)
    add_para(doc, "Доступ к административным данным разрешается только пользователям с действительной обычной сессией, разрешенным email для входа в /admin и отдельным admin_session JWT claim после второго admin-panel login. Обычная роль пользователя не должна быть достаточным основанием для доступа к admin API.")
    add_heading(doc, "5.9. Сохранность информации при авариях", 2)
    add_bullets(doc, ["Регулярные резервные копии PostgreSQL с проверкой восстановления.", "Отсутствие хранения пользовательских файлов в текущем scope.", "Документированный порядок восстановления после сбоя БД, API и web.", "Журналы аудита и login attempts должны сохраняться с учетом требований хранения и защиты."])
    add_heading(doc, "5.10. Защита от внешних воздействий", 2)
    add_para(doc, "Физическая и сетевая защита обеспечивается инфраструктурой. На уровне приложения требуются ограничения trusted hosts, CORS, size limits, rate limiting, безопасные cookie flags, HSTS на production ingress и исключение публичного доступа к БД.")
    add_heading(doc, "5.11. Патентная чистота", 2)
    add_para(doc, "Используемые open-source компоненты должны применяться в соответствии с их лицензиями. Исполнитель обязан вести перечень зависимостей и не включать компоненты с лицензиями, несовместимыми с условиями проекта.")
    add_heading(doc, "5.12. Стандартизация и унификация", 2)
    add_para(doc, "Должны использоваться единые соглашения по TypeScript/Python-коду, API contracts, миграциям, валидации, именованию статусов, двуязычным материалам и компонентам интерфейса. Все active site sections должны использовать общие схемы локализованного контента.")
    doc.add_page_break()

    section_title(doc, 6, "Требования к функциям (задачам)")
    add_heading(doc, "6.1. Перечень функций по подсистемам", 2)
    add_bullets(doc, ["Публичный контент: чтение новостей, страниц, нормативной базы и контактной информации региональных подразделений.", "Обращения: создание обращения, выдача tracking code, lookup статуса, административная смена статуса.", "Аутентификация: Telegram start/webhook/complete, phone confirmation, email/password register/login, refresh, logout, /auth/me.", "Кабинет кандидата: отображение профиля, заявки, статуса, результатов психологических тестов.", "Тестирование: прохождение primary-selection, сохранение результата, чтение собственных результатов и админский просмотр.", "Администрирование: dashboard, списки, фильтрация, модерация статусов, content CRUD, audit logging."])
    add_heading(doc, "6.2. Очередность ввода функций", 2)
    add_matrix_table(doc, ["Очередь", "Состав", "Условие готовности"], [["1", "Базовые public/auth/appeals/candidate/admin flows.", "Сквозные сценарии работают на ru/kk, критические ошибки устранены."], ["2", "Content CRUD для активных разделов и расширенный аудит.", "Администратор управляет контентом без прямого доступа к БД."], ["3", "Security hardening и эксплуатационная готовность.", "Пройдены проверки конфигурации, тесты, build, backup/restore dry run."], ["4", "Дальнейшие функции по решению Заказчика.", "Оформлено отдельное изменение scope и ТЗ/допсоглашение."]], [900, 5200, 3260])
    add_heading(doc, "6.3. Временной регламент реализации функций", 2)
    add_para(doc, "Публичные чтения и кабинетные операции должны выполняться интерактивно без необоснованных задержек. Для типовых API-запросов целевое время ответа на application-level при нормальной нагрузке - до 1 секунды для чтения и до 2 секунд для записи, без учета внешней сети и инфраструктурных задержек.")
    add_heading(doc, "6.4. Форма выходной информации и качество реализации", 2)
    add_bullets(doc, ["API возвращает JSON с предсказуемыми схемами и безопасными ошибками.", "Frontend отображает локализованные сообщения на русском и казахском языках.", "Статусы должны иметь человекочитаемые labels и машинные enum-значения.", "Административные списки должны иметь clear empty/error/loading states.", "Не допускается раскрытие stack trace, секретов, SQL и внутренней структуры в публичном ответе."])
    add_heading(doc, "6.5. Отказы по функциям", 2)
    add_matrix_table(doc, ["Функция", "Критерий отказа", "Ожидаемая реакция"], [["Подача обращения", "Данные не сохранены или не выдан tracking code.", "Показать ошибку, записать server log, не создавать частичный пользовательский успех."], ["Refresh session", "Токен недействителен, истек или отозван.", "Очистить frontend session state и запросить повторный вход."], ["Admin action", "Нет admin_session или прав.", "Вернуть 401/403, не менять данные, записать security-relevant event где применимо."], ["Тестирование", "Результат не сохранен после завершения.", "Показать ошибку сохранения и позволить повторить отправку без потери локального результата."]], [1900, 3700, 3760])
    doc.add_page_break()

    section_title(doc, 7, "Требования к видам обеспечения")
    add_heading(doc, "7.1. Математическое обеспечение", 2)
    add_para(doc, "Сложное математическое обеспечение на текущем этапе не требуется. Психологическое тестирование сохраняет ответы, временные показатели и summaries; автоматическая оценка и интерпретация отключены до передачи утвержденных ключей и методики Заказчиком.")
    add_heading(doc, "7.2. Информационное обеспечение", 2)
    add_para(doc, "Данные хранятся в PostgreSQL. Основные сущности: User, CandidateApplication, Appeal, RefreshSession, LoginAttempt, TelegramLoginChallenge, PsychologicalTestResult, PsychologicalTestProgress, RegionOffice, AuditLog, News, Page. Структура данных должна поддерживаться миграциями Alembic.")
    add_heading(doc, "7.3. Лингвистическое обеспечение", 2)
    add_para(doc, "Поддерживаются только русский и казахский языки. Все публичные маршруты должны быть locale-prefixed: /ru и /kk. Тексты интерфейса, статусы, ошибки и основные материалы должны иметь согласованные переводы; английская локаль не реализуется.")
    add_heading(doc, "7.4. Программное обеспечение", 2)
    add_bullets(doc, ["Frontend: Next.js App Router, React, TypeScript, Tailwind CSS, lucide-react, Leaflet/react-leaflet.", "Backend: Python FastAPI, SQLAlchemy 2, Pydantic, Alembic.", "DB: PostgreSQL для Docker/production, SQLite только для локальной разработки.", "Инструменты качества: frontend typecheck/build, backend tests, миграционные проверки, dependency audit."])
    add_heading(doc, "7.5. Техническое обеспечение", 2)
    add_para(doc, "Минимальная production-схема: web container, API container, PostgreSQL, reverse proxy/ingress с TLS, WAF/rate limiting на периметре, закрытая сеть БД, хранилище резервных копий, логирование и мониторинг. Конкретные мощности определяются нагрузочным профилем и требованиями hosting party.")
    add_heading(doc, "7.6. Метрологическое обеспечение", 2)
    add_para(doc, "Специальное метрологическое обеспечение не требуется. Для performance и availability должны использоваться измеримые метрики: время ответа API, доля ошибок, uptime, нагрузка БД, заполнение дисков, успешность backup/restore.")
    add_heading(doc, "7.7. Организационное, методическое и другие виды обеспечения", 2)
    add_bullets(doc, ["Регламент публикации и модерации официальной информации.", "Регламент обработки обращений и кандидатских заявок.", "Инструкция администратора и модератора.", "Инструкция по deployment, rollback, backup/restore и работе с инцидентами.", "Матрица ролей и прав доступа."])
    doc.add_page_break()

    section_title(doc, 8, "Состав и содержание работ по созданию системы")
    add_heading(doc, "8.1. Стадии и этапы работ", 2)
    add_matrix_table(doc, ["Этап", "Содержание", "Результат"], [["1. Обследование и уточнение требований", "Согласование scope, ролей, данных, ограничений и приемочных критериев.", "Утвержденное ТЗ и backlog."], ["2. Проектирование", "API contracts, data model, security model, UX flows, deployment outline.", "Проектные решения и план миграций."], ["3. Разработка", "Frontend, backend, migrations, admin/content CRUD, тесты.", "Рабочая версия в staging."], ["4. Испытания", "Functional, regression, security hardening checks, build/typecheck.", "Протокол испытаний и defects list."], ["5. Внедрение", "Production configuration, migrations, seed/reference data, monitoring.", "Готовая к эксплуатации версия."], ["6. Сопровождение", "Гарантийные исправления, консультации, minor hardening.", "Акты выполненных работ и release notes."]], [1200, 4900, 3260])
    add_heading(doc, "8.2. Документы по окончании стадий", 2)
    add_bullets(doc, ["Техническое задание и приложения.", "Архитектурное описание.", "Описание API и схем данных.", "Инструкция администратора/модератора.", "Инструкция по развертыванию и эксплуатации.", "Программа и методика испытаний.", "Протоколы испытаний и акт приемки.", "Перечень переменных окружения без реальных секретов."])
    add_heading(doc, "8.3. Экспертиза технической документации", 2)
    add_para(doc, "Экспертиза проводится Заказчиком, уполномоченной hosting/reviewing party и при необходимости специализированным подразделением информационной безопасности. Проверяются полнота требований, соответствие scope, security posture, эксплуатационная готовность и непротиворечивость документации фактической реализации.")
    add_heading(doc, "8.4. Работы по обеспечению надежности", 2)
    add_bullets(doc, ["Разработка backend и frontend тестов для критических flows.", "Проверка миграций на чистой и существующей БД.", "Настройка healthchecks, logs, monitoring и backup/restore.", "Контроль secret/config validation на production startup.", "Проверка graceful handling ошибок API и недоступности сервисов."])
    add_heading(doc, "8.5. Работы по метрологическому обеспечению", 2)
    add_para(doc, "Не требуются в классическом метрологическом смысле. Вместо этого фиксируются эксплуатационные метрики и пороги мониторинга, подлежащие согласованию перед промышленным запуском.")
    doc.add_page_break()

    section_title(doc, 9, "Порядок контроля и приемки системы")
    add_heading(doc, "9.1. Верификация", 2)
    add_bullets(doc, ["Непротиворечивость требований и соответствие scope.", "Работоспособность ru/kk маршрутов.", "Соответствие API schemas и frontend expectations.", "Корректность auth/session flows, включая refresh rotation и logout.", "Проверка RBAC/admin_session separation.", "Проверка отсутствия out-of-scope функций.", "Наличие миграций Alembic для schema changes.", "Прохождение тестов, typecheck, build и smoke checks."])
    add_heading(doc, "9.2. Аттестация", 2)
    add_para(doc, "Аттестация означает подтверждение соответствия системы настоящему ТЗ, утвержденным проектным решениям, документации и требованиям безопасного применения. Итоговая аттестация проводится после устранения критических и высоких дефектов.")
    add_heading(doc, "9.3. Совместная проверка", 2)
    add_para(doc, "Совместная проверка проводится представителями Заказчика, Исполнителя и при необходимости hosting/reviewing party. Проверяются функциональные сценарии, настройки окружения, резервное копирование, права доступа и эксплуатационная документация.")
    add_heading(doc, "9.4. Аудит", 2)
    add_para(doc, "Аудит устанавливает соответствие фактических работ ТЗ, плану и договору. Отдельно проверяются административные журналы, login attempts, действия с обращениями, кандидатскими заявками и контентом.")
    add_heading(doc, "9.5. Разрешение проблем", 2)
    add_para(doc, "Каждое обнаруженное несоответствие должно быть идентифицировано, описано, классифицировано по критичности, назначено ответственному, исправлено и повторно проверено. Критические дефекты безопасности и потери данных блокируют приемку.")
    add_heading(doc, "9.6. Приемка работ по стадиям", 2)
    add_matrix_table(doc, ["Стадия", "Критерий приемки"], [["ТЗ", "Документ согласован Заказчиком, scope и ограничения зафиксированы."], ["Разработка", "Функции реализованы, миграции есть, демо-учетки и hardcoded credentials отсутствуют."], ["Испытания", "Критические flows проверены, defects list закрыт по agreed threshold."], ["Внедрение", "Production configuration подготовлена, backup/restore и healthchecks проверены."], ["Сопровождение", "Гарантийные дефекты устранены в согласованные сроки."]], [1800, 7560])
    doc.add_page_break()

    section_title(doc, 10, "Подготовка объекта автоматизации к вводу системы в действие")
    add_heading(doc, "10.1. Приведение информации к виду, пригодному для обработки", 2)
    add_para(doc, "Контент и справочные данные должны быть подготовлены в двуязычной структуре ru/kk, очищены от черновых материалов, проверены на официальность формулировок и загружены через предусмотренные seed/import/admin механизмы без прямого ручного изменения production БД.")
    add_heading(doc, "10.2. Изменения в объекте автоматизации", 2)
    add_bullets(doc, ["Назначить ответственных администраторов и модераторов.", "Утвердить регламент публикации и обработки обращений.", "Подготовить production secrets и переменные окружения.", "Настроить домены, TLS, WAF/rate limiting, сети, backup, logs и monitoring.", "Обеспечить процесс согласования релизов и отката."])
    add_heading(doc, "10.3. Создание условий функционирования", 2)
    add_para(doc, "До запуска должны быть подготовлены staging и production окружения, закрытый доступ к БД, secure cookie flags, production secret checks, CORS/trusted hosts allowlists, миграционная процедура и проверенный plan восстановления.")
    add_heading(doc, "10.4. Создание подразделений и служб", 2)
    add_para(doc, "Специальное новое подразделение не требуется, однако должны быть назначены владельцы процессов: контент, обработка обращений, работа с кандидатами, техническая эксплуатация, информационная безопасность и приемка изменений.")
    add_heading(doc, "10.5. Комплектование штатов и обучение", 2)
    add_para(doc, "Перед вводом системы администраторы и модераторы проходят инструктаж по входу в административную панель, обработке записей, работе со статусами, правилам публикации материалов, требованиям по защите персональных данных и порядку обращения с инцидентами.")
    doc.add_page_break()

    section_title(doc, 11, "Требования к документированию")
    add_heading(doc, "11.1. Перечень комплектов документов", 2)
    add_bullets(doc, ["Техническое задание.", "Архитектурное описание системы.", "Описание API endpoints и схем данных.", "Описание модели ролей и прав.", "Инструкция администратора и модератора.", "Инструкция пользователя-кандидата при необходимости.", "Руководство по развертыванию, миграциям и эксплуатации.", "Программа и методика испытаний.", "Release notes и changelog.", "Security hardening checklist и перечень production env variables."])
    add_heading(doc, "11.2. Документы на машинных носителях", 2)
    add_para(doc, "Документы предоставляются в форматах DOCX/PDF/Markdown по согласованию. Исходный код, миграции, конфигурационные шаблоны, тесты и документация передаются через репозиторий проекта или иной утвержденный канал.")
    add_heading(doc, "11.3. Документирование комплектующих элементов", 2)
    add_para(doc, "Документирование open-source зависимостей включает список backend/frontend packages, версий, лицензий и известных ограничений. Компоненты общего назначения описываются в эксплуатационной документации в части, необходимой для сопровождения.")
    add_heading(doc, "11.4. Дополнительные требования к составу документов", 2)
    add_bullets(doc, ["Документация не должна содержать реальные production-секреты.", "Команды запуска должны быть воспроизводимыми для локального, Docker и production-подобного окружения.", "Все документы должны соответствовать фактическому состоянию исходного кода на момент передачи.", "Разделы безопасности должны отделять application-level controls от инфраструктурных предпосылок."])
    doc.add_page_break()

    section_title(doc, 12, "Источники разработки")
    add_heading(doc, "12.1. Документы и информационные материалы", 2)
    add_bullets(doc, ["Шаблон технического задания «техническое_задание_2026-07-13.docx».", "Репозиторий проекта официального интернет-портала КНБ РК.", "AGENTS.md с проектным контекстом, scope, текущим состоянием и приоритетами.", "docs/ARCHITECTURE.md.", "SECURITY_ROADMAP.md.", "docker-compose.yml, Dockerfiles, backend/frontend configuration.", "Исходный код apps/web и apps/api, включая API routes, модели, миграции, компоненты пользовательского кабинета и admin panel.", "Применимые нормативные и организационные документы Заказчика, подлежащие уточнению на этапе согласования."])
    add_para(doc, "Настоящее ТЗ должно актуализироваться при изменении функционального scope, архитектуры, требований безопасности, интеграционных решений или порядка эксплуатации системы.", bold=True)

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT.resolve())
